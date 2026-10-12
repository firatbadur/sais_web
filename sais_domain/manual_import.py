"""
Manuel veri girişi motoru — analizör SD kartından alınan verinin Excel ile içe aktarımı.

Elektrik kesintisinde polling durur; historian'da ve Bakanlık'ta (SİM) boşluk
oluşur. Kayıp dakikaların değerleri analizörün SD kartından alınıp **istasyona
özel Excel şablonuna** girilir, sayfadan yüklenir.

Akış (tek doğruluk kaynağı bu modül):

1. :func:`build_template` — kabinin Bakanlık'a gönderilen analog kolonlarıyla
   şablon (.xlsx): ``Veri`` (Tarih Saat | Durum | <parametre kodları>),
   ``Açıklama`` (kurallar + örnek satırlar), ``Durum Kodları`` (Bakanlık kataloğu).
2. :func:`stage_upload` → :func:`parse_workbook` — dosyayı doğrular, normalize
   eder; hatalar varsa onay verilemez. Sonuç ``ManualDataImport(staged)``.
3. :func:`apply_import` (Celery ``apply_manual_import``) — ``Reading``'e
   ``origin="manual"`` ile yazar (**o dakikada sensörün okuması varsa ATLANIR**),
   aggregate'leri yeniden hesaplar, kabul penceresindeki dakikaları SİM gönderim
   kuyruğuna (backfill şeridi, ``SOURCE_MANUAL``) ekler.

**Status satır başına zorunludur** — yıkama/bakım/arıza dakikaları "geçerli"
gibi gönderilmesin. ``persist_reading`` bilinçli olarak KULLANILMAZ: canlı
yıkama overlay'i + throttle uygular ve ``SensorLatest``'i (HMI anlık değeri)
eski veriyle ezerdi.
"""
from __future__ import annotations

import io
import re
from collections import Counter, defaultdict
from datetime import datetime, timedelta

from django.db import transaction
from django.utils import timezone

from .sim_report import SIM_DATA_STATUS_CODES, STATUS_BY_CODE, VALID_CODES

DATA_SHEET = "Veri"
HELP_SHEET = "Açıklama"
CODES_SHEET = "Durum Kodları"
DATE_HEADER = "Tarih Saat"
STATUS_HEADER = "Durum"
DATE_ALIASES = {"tarih saat", "tarihsaat", "tarih", "readtime", "zaman"}
STATUS_ALIASES = {"durum", "status", "durum kodu"}

# Şablon kimliği (workbook keywords) — başka kabinin şablonu yüklenirse uyarı.
TEMPLATE_KEYWORD_PREFIX = "envisoft_manual_import:"

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20000          # ~14 gün dakikalık veri
MAX_LISTED_ISSUES = 200   # ekranda listelenen hata/uyarı tavanı
PREVIEW_ROWS = 50
TEMPLATE_FORMATTED_ROWS = 1440  # tarih biçimi + açılır liste uygulanan satır (1 gün)

DATE_FORMATS = (
    "%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M",
    "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M",
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M",
    "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M",
)

# Yıkama/bakım/operasyonel durumlar → "uncertain" (persist_reading overlay'iyle
# aynı); diğer geçersiz kodlar → "bad". İkisi de aggregate ortalamasına girmez.
_UNCERTAIN_CODES = {23, 24, 25, 26, 30, 39} | set(VALID_CODES)


class ManualImportError(Exception):
    """Dosya hiç işlenemedi (boyut, biçim, kabin yok…)."""


def status_quality(code: int) -> str:
    if code == 1:
        return "good"
    if code in _UNCERTAIN_CODES:
        return "uncertain"
    return "bad"


# ---------------------------------------------------------------------------
# Kolonlar
# ---------------------------------------------------------------------------
def template_columns(cabinet, *, include_unsent: bool = False) -> list[dict]:
    """Kabin için kolonlar.

    İstasyondaki aktif **analog** sensörlerin parametreleri. Varsayılan
    (şablon) yalnız Bakanlık'a gönderilenleri döndürür — ``SimChannel`` seçimi
    yapılmışsa ``build_sim_payload`` ile aynı kural. ``include_unsent=True``
    (doğrulama/uygulama) gönderilmeyen kolonları da ``sent=False`` ile ekler:
    kolon seçimi şablon indirildikten SONRA değişirse dosya yine kabul edilir,
    o kolonlar yalnız historian'a yazılır (SİM payload'u zaten
    ``build_sim_payloads_for_times`` içinde ayrıca süzülür).
    Aynı parametreye bağlı birden çok sensör varsa ilki (en küçük id) kullanılır.
    """
    from api.models import Sensor

    from .models import SimChannel

    izinli = SimChannel.enabled_parameter_names(cabinet)
    sensors = (
        Sensor.objects
        .filter(
            connection__station_id=cabinet.station_id,
            is_active=True,
            parameter__isnull=False,
            sensor_type__in=(0, 1),
        )
        .select_related("parameter")
        .order_by("parameter__parameter_name", "id")
    )
    cols: dict[str, dict] = {}
    for sn in sensors:
        p = sn.parameter
        code = (p.parameter_name or p.parameter_txt or "").strip()
        if not code or code in cols:
            continue
        sent = izinli is None or code in izinli
        if not sent and not include_unsent:
            continue
        cols[code] = {
            "code": code,
            "sent": sent,
            "text": (p.parameter_txt or p.parameter_name or "").strip(),
            "unit": (p.unit or "").strip(),
            "sensor": sn,
            "min": p.min_range,
            "max": p.max_range,
        }
    return list(cols.values())


# ---------------------------------------------------------------------------
# Şablon
# ---------------------------------------------------------------------------
def build_template(cabinet) -> bytes:
    """Kabine özel boş şablon (.xlsx) üretir."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    cols = template_columns(cabinet)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1B84FF")
    hint_font = Font(italic=True, color="808080", size=9)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb = Workbook()
    wb.properties.title = f"Manuel Veri Girişi — {cabinet}"
    wb.properties.keywords = f"{TEMPLATE_KEYWORD_PREFIX}{cabinet.pk}"

    # --- Veri ---
    ws = wb.active
    ws.title = DATA_SHEET
    headers = [DATE_HEADER, STATUS_HEADER] + [c["code"] for c in cols]
    hints = ["(gg.aa.yyyy ss:dd)", "(kod — listeden)"] + [
        f"({c['text']}{', ' + c['unit'] if c['unit'] else ''})" for c in cols
    ]
    for i, (h, hint) in enumerate(zip(headers, hints), start=1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font, cell.fill, cell.alignment = head_font, head_fill, center
        hc = ws.cell(row=2, column=i, value=hint)
        hc.font, hc.alignment = hint_font, center
        ws.column_dimensions[get_column_letter(i)].width = 20 if i == 1 else 16
    ws.freeze_panes = "C3"
    for r in range(3, 3 + TEMPLATE_FORMATTED_ROWS):
        ws.cell(row=r, column=1).number_format = "dd.mm.yyyy hh:mm"

    dv = DataValidation(
        type="list",
        formula1=f"'{CODES_SHEET}'!$A$2:$A${len(SIM_DATA_STATUS_CODES) + 1}",
        allow_blank=False,
        showErrorMessage=True,
        errorTitle="Geçersiz durum kodu",
        error="Durum kodu 'Durum Kodları' sayfasındaki listeden seçilmelidir.",
        promptTitle="Durum kodu",
        prompt="1 = Veri Geçerli, 23 = Yıkama, 24 = Haftalık Yıkama, 25 = İstasyon Bakımda ...",
        showInputMessage=True,
    )
    ws.add_data_validation(dv)
    dv.add(f"B3:B{2 + MAX_ROWS}")

    # --- Açıklama ---
    wh = wb.create_sheet(HELP_SHEET)
    wh.column_dimensions["A"].width = 22
    wh.column_dimensions["B"].width = 14
    lines = [
        (f"Manuel Veri Girişi — {cabinet}", True),
        ("Analizörün SD kartından alınan veriler 'Veri' sayfasına girilir.", False),
        ("", False),
        ("Kurallar", True),
        ("• Her satır bir dakikadır. 'Tarih Saat' gg.aa.yyyy ss:dd biçiminde (saniye yok sayılır).", False),
        ("• 'Durum' ZORUNLUDUR ve satırdaki tüm parametrelere uygulanır (1 = Veri Geçerli).", False),
        ("  Yıkama = 23, Haftalık Yıkama = 24, İstasyon Bakımda = 25, Tesis Bakımda = 26 ...", False),
        ("  Tüm kodlar 'Durum Kodları' sayfasındadır.", False),
        ("• Değer ondalık ayracı virgül veya nokta olabilir. Boş hücre = o parametre girilmedi.", False),
        ("• Sütun başlıkları (1. satır) değiştirilmemelidir; 2. satır açıklamadır, silinebilir.", False),
        ("• Sistemde o dakika için zaten okuma varsa o değer ATLANIR (mevcut veri korunur).", False),
        ("• Bakanlık kabul penceresinden (varsayılan 48 saat) eski dakikalar yalnız sistem", False),
        ("  geçmişine yazılır, Bakanlık'a gönderilemez.", False),
        (f"• En fazla {MAX_ROWS} satır yüklenebilir.", False),
        ("", False),
        ("Örnek (bu sayfadaki satırlar içe aktarılmaz)", True),
    ]
    for r, (text, bold) in enumerate(lines, start=1):
        c = wh.cell(row=r, column=1, value=text)
        if bold:
            c.font = Font(bold=True, size=12 if r == 1 else 11)
    ex_row = len(lines) + 1
    for i, h in enumerate(headers, start=1):
        cell = wh.cell(row=ex_row, column=i, value=h)
        cell.font, cell.fill, cell.alignment = head_font, head_fill, center
    base = timezone.localtime().replace(hour=9, minute=0, second=0, microsecond=0, tzinfo=None)
    example_status = [1, 1, 23, 23, 1]
    for k, code in enumerate(example_status):
        rr = ex_row + 1 + k
        dc = wh.cell(row=rr, column=1, value=base + timedelta(minutes=k))
        dc.number_format = "dd.mm.yyyy hh:mm"
        wh.cell(row=rr, column=2, value=code)
        for j, c in enumerate(cols, start=3):
            lo = c["min"] if c["min"] is not None else 0
            hi = c["max"] if c["max"] is not None else 100
            wh.cell(row=rr, column=j, value=round(lo + (hi - lo) * (0.3 + 0.02 * k), 2))

    # --- Durum Kodları ---
    wc = wb.create_sheet(CODES_SHEET)
    for i, h in enumerate(("Kod", "Ad", "Açıklama", "Geçerli Veri"), start=1):
        cell = wc.cell(row=1, column=i, value=h)
        cell.font, cell.fill = head_font, head_fill
    for r, s in enumerate(SIM_DATA_STATUS_CODES, start=2):
        wc.cell(row=r, column=1, value=s["code"])
        wc.cell(row=r, column=2, value=s["name"])
        wc.cell(row=r, column=3, value=s["desc"])
        wc.cell(row=r, column=4, value="Evet" if s["valid"] else "Hayır")
    for col, w in zip("ABCD", (8, 30, 40, 12)):
        wc.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Doğrulama
# ---------------------------------------------------------------------------
def _parse_dt(raw):
    """Hücre → TZ-aware dakika (saniyesiz) veya None."""
    from openpyxl.utils.datetime import from_excel

    if raw is None or raw == "":
        return None
    if isinstance(raw, datetime):
        dt = raw
    elif isinstance(raw, (int, float)) and not isinstance(raw, bool):
        try:
            dt = from_excel(raw)
        except (ValueError, TypeError, OverflowError):
            return None
        if not isinstance(dt, datetime):
            return None
    elif isinstance(raw, str):
        text = raw.strip()
        dt = None
        for fmt in DATE_FORMATS:
            try:
                dt = datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
        if dt is None:
            return None
    else:
        return None
    if timezone.is_naive(dt):
        dt = timezone.make_aware(dt)
    return dt.replace(second=0, microsecond=0)


_LEADING_INT = re.compile(r"^\s*(-?\d+)")


def _parse_status(raw):
    if raw is None or raw == "":
        return None
    if isinstance(raw, bool):
        return None
    if isinstance(raw, int):
        return raw
    if isinstance(raw, float):
        return int(raw) if raw.is_integer() else None
    m = _LEADING_INT.match(str(raw))  # "23", "23 - Yıkama"
    return int(m.group(1)) if m else None


def _parse_value(raw):
    """(değer, geçerli_mi). Boş → (None, True)."""
    if raw is None:
        return None, True
    if isinstance(raw, bool):
        return None, False
    if isinstance(raw, (int, float)):
        return float(raw), True
    text = str(raw).strip()
    if not text:
        return None, True
    text = text.replace(" ", "")
    if "," in text and "." in text:      # 1.234,5 → 1234.5
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", ".")
    try:
        return float(text), True
    except ValueError:
        return None, False


def _fmt(dt) -> str:
    return timezone.localtime(dt).strftime("%d.%m.%Y %H:%M")


def parse_workbook(cabinet, data: bytes) -> dict:
    """Yüklenen dosyayı doğrular ve normalize eder.

    Dönen sözlük: ``rows`` (normalize satırlar), ``errors``, ``warnings``,
    ``summary``, ``range_start``, ``range_end``. Hata listesi boş değilse
    içe aktarım onaylanamaz.
    """
    from openpyxl import load_workbook

    from .sim_outbox import accept_window_hours

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — bozuk/yanlış dosya türü
        raise ManualImportError(
            f"Dosya Excel (.xlsx) olarak okunamadı: {type(exc).__name__}"
        ) from exc

    errors: list[str] = []
    warnings: list[str] = []
    error_total = 0

    def err(msg):
        nonlocal error_total
        error_total += 1
        if len(errors) < MAX_LISTED_ISSUES:
            errors.append(msg)

    keywords = (wb.properties.keywords or "") if wb.properties else ""
    if keywords.startswith(TEMPLATE_KEYWORD_PREFIX):
        if keywords[len(TEMPLATE_KEYWORD_PREFIX):] != str(cabinet.pk):
            warnings.append(
                "Bu dosya başka bir kabin için indirilen şablondan üretilmiş görünüyor; "
                "kolonların bu kabinle eşleştiğini kontrol edin."
            )

    ws = wb[DATA_SHEET] if DATA_SHEET in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = next(it, None)
    if not header:
        wb.close()
        raise ManualImportError("Dosyada başlık satırı yok.")

    cols = {c["code"]: c for c in template_columns(cabinet, include_unsent=True)}
    cols_ci = {k.lower(): k for k in cols}
    date_idx = status_idx = None
    value_idx: dict[int, str] = {}
    for i, h in enumerate(header):
        name = str(h).strip() if h is not None else ""
        if not name:
            continue
        low = name.lower()
        if low in DATE_ALIASES and date_idx is None:
            date_idx = i
        elif low in STATUS_ALIASES and status_idx is None:
            status_idx = i
        elif low in cols_ci:
            code = cols_ci[low]
            if code in value_idx.values():
                err(f"'{name}' sütunu birden fazla kez var.")
            else:
                value_idx[i] = code
        else:
            err(
                f"Bilinmeyen sütun: '{name}'. Bu kabinin kolonları: "
                f"{', '.join(cols) or '—'}."
            )
    if date_idx is None:
        err(f"'{DATE_HEADER}' sütunu bulunamadı.")
    if status_idx is None:
        err(f"'{STATUS_HEADER}' sütunu bulunamadı.")
    if not value_idx:
        err("Hiç parametre sütunu bulunamadı.")
    if date_idx is None or status_idx is None or not value_idx:
        wb.close()
        return {
            "rows": [], "errors": errors, "warnings": warnings,
            "summary": {"error_total": error_total}, "range_start": None, "range_end": None,
        }

    unsent_cols = [c for c in value_idx.values() if not cols[c]["sent"]]
    if unsent_cols:
        warnings.append(
            f"Şu kolonlar Bakanlık'a gönderilmiyor (Gönderilecek Veriler sayfası); "
            f"yalnız sistem geçmişine yazılacak: {', '.join(unsent_cols)}."
        )
    missing_cols = [
        c for c, col in cols.items() if col["sent"] and c not in value_idx.values()
    ]
    if missing_cols:
        warnings.append(
            f"Şablondaki şu kolonlar dosyada yok, bu parametreler girilmeyecek: "
            f"{', '.join(missing_cols)}."
        )

    now = timezone.now()
    rows: list[dict] = []
    seen: dict[datetime, int] = {}
    out_of_range: Counter = Counter()
    status_counts: Counter = Counter()
    value_count = 0
    too_many = False

    for excel_row, rec in enumerate(it, start=2):
        rec = list(rec or ())
        cells = rec + [None] * (len(header) - len(rec))
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in cells):
            continue
        raw_date = cells[date_idx]
        # Şablonun 2. satırı açıklama satırıdır: "(gg.aa.yyyy ss:dd)"
        if isinstance(raw_date, str) and raw_date.strip().startswith("("):
            continue
        if len(rows) >= MAX_ROWS:
            too_many = True
            break

        dt = _parse_dt(raw_date)
        if dt is None:
            err(f"Satır {excel_row}: tarih okunamadı ({raw_date!r}). Biçim: gg.aa.yyyy ss:dd.")
            continue
        if dt > now:
            err(f"Satır {excel_row}: tarih gelecekte ({_fmt(dt)}).")
            continue
        if dt in seen:
            err(f"Satır {excel_row}: {_fmt(dt)} dakikası {seen[dt]}. satırda da var.")
            continue
        seen[dt] = excel_row

        raw_status = cells[status_idx]
        code = _parse_status(raw_status)
        if code is None:
            err(f"Satır {excel_row}: Durum boş veya sayı değil ({raw_status!r}).")
            continue
        if code not in STATUS_BY_CODE:
            err(f"Satır {excel_row}: {code} geçerli bir Bakanlık durum kodu değil.")
            continue

        values: dict[str, float] = {}
        row_ok = True
        for i, pcode in value_idx.items():
            val, ok = _parse_value(cells[i])
            if not ok:
                err(f"Satır {excel_row}, {pcode}: sayı değil ({cells[i]!r}).")
                row_ok = False
                continue
            if val is None:
                continue
            col = cols[pcode]
            lo, hi = col["min"], col["max"]
            if (lo is not None and val < lo) or (hi is not None and val > hi):
                out_of_range[pcode] += 1
            values[pcode] = val
        if not row_ok:
            continue
        if not values:
            err(f"Satır {excel_row}: hiç değer girilmemiş.")
            continue
        status_counts[code] += 1
        value_count += len(values)
        rows.append({"t": dt.isoformat(), "s": code, "v": values})

    wb.close()

    if too_many:
        err(f"Dosyada {MAX_ROWS} satırdan fazla veri var; dosyayı bölerek yükleyin.")
    if not rows and not error_total:
        err("Dosyada içe aktarılacak veri satırı yok.")

    for pcode, n in out_of_range.items():
        c = cols[pcode]
        warnings.append(
            f"{pcode}: {n} değer ölçüm aralığı dışında ({c['min']} – {c['max']})."
        )

    range_start = range_end = None
    existing = outside_window = 0
    if rows:
        times = [datetime.fromisoformat(r["t"]) for r in rows]
        range_start, range_end = min(times), max(times)
        existing_keys = _existing_keys(
            [c["sensor"].id for c in cols.values()], range_start, range_end,
        )
        for r, t in zip(rows, times):
            for pcode in r["v"]:
                if (cols[pcode]["sensor"].id, t) in existing_keys:
                    existing += 1
        if existing:
            warnings.append(
                f"{existing} değer için sistemde o dakikada zaten okuma var; "
                f"bunlar ATLANACAK (mevcut veri korunur)."
            )
        window_h = accept_window_hours()
        cutoff = now - timedelta(hours=window_h)
        outside_window = sum(1 for t in times if t < cutoff)
        if outside_window:
            warnings.append(
                f"{outside_window} dakika Bakanlık kabul penceresinin ({window_h} saat) "
                f"dışında: yalnız sistem geçmişine yazılacak, Bakanlık'a gönderilmeyecek. "
                f"Pencere Sistem Kontrol sayfasından geçici olarak genişletilebilir."
            )

    summary = {
        "row_count": len(rows),
        "value_count": value_count,
        "columns": list(value_idx.values()),
        "status_counts": {str(k): v for k, v in sorted(status_counts.items())},
        "status_labels": {
            str(k): STATUS_BY_CODE[k]["desc"] for k in status_counts
        },
        "existing_count": existing,
        "outside_window_count": outside_window,
        "error_total": error_total,
        "preview": [
            {
                "t": _fmt(datetime.fromisoformat(r["t"])),
                "s": r["s"],
                "valid": r["s"] in VALID_CODES,
                "v": [r["v"].get(c) for c in value_idx.values()],
            }
            for r in rows[:PREVIEW_ROWS]
        ],
    }
    return {
        "rows": rows, "errors": errors, "warnings": warnings, "summary": summary,
        "range_start": range_start, "range_end": range_end,
    }


def _existing_keys(sensor_ids, start, end) -> set[tuple[int, datetime]]:
    """Aralıkta okuması olan (sensör, dakika) çiftleri."""
    from api.models import Reading

    keys = set()
    qs = (
        Reading.objects
        .filter(
            sensor_id__in=sensor_ids,
            time_iso__gte=start,
            time_iso__lt=end + timedelta(minutes=1),
        )
        .values_list("sensor_id", "time_iso")
    )
    for sid, ts in qs.iterator():
        keys.add((sid, ts.replace(second=0, microsecond=0)))
    return keys


def stage_upload(cabinet, data: bytes, filename: str, user=None):
    """Dosyayı doğrulayıp ``ManualDataImport(staged)`` olarak kaydeder.

    Aynı kabinin önceki onay bekleyen yüklemesi iptal edilir (tek bekleyen).
    """
    from .models import ManualDataImport

    if len(data) > MAX_UPLOAD_BYTES:
        raise ManualImportError(
            f"Dosya çok büyük ({len(data) // 1024} KB). En fazla "
            f"{MAX_UPLOAD_BYTES // (1024 * 1024)} MB yüklenebilir."
        )
    if not (filename or "").lower().endswith(".xlsx"):
        raise ManualImportError("Yalnız Excel (.xlsx) dosyası yüklenebilir.")

    result = parse_workbook(cabinet, data)
    ManualDataImport.objects.filter(
        cabinet=cabinet, status=ManualDataImport.STATUS_STAGED,
    ).update(status=ManualDataImport.STATUS_DISCARDED, rows=[])
    return ManualDataImport.objects.create(
        cabinet=cabinet,
        user=user if (user and user.is_authenticated) else None,
        filename=(filename or "")[:255],
        rows=result["rows"] if not result["errors"] else [],
        errors=result["errors"],
        warnings=result["warnings"],
        summary=result["summary"],
        range_start=result["range_start"],
        range_end=result["range_end"],
    )


# ---------------------------------------------------------------------------
# Uygulama
# ---------------------------------------------------------------------------
def _status_obj(code: int):
    from api.models import StatusCode

    obj = StatusCode.objects.filter(code=code).first()
    if obj is None:
        meta = STATUS_BY_CODE.get(code, {})
        obj = StatusCode.objects.create(code=code, name=meta.get("desc") or str(code))
    return obj


def apply_import(record) -> dict:
    """Onaylanmış içe aktarımı uygular (Celery task'ından çağrılır).

    1. Reading insert (``origin="manual"``) — o dakikada sensörün okuması
       varsa değer ATLANIR. Insert + aggregate tek transaction.
    2. Etkilenen aralığın aggregate'leri yeniden hesaplanır.
    3. Commit sonrası: en az bir değer eklenen ve kabul penceresindeki
       dakikalar SİM kuyruğuna (backfill şeridi, ``SOURCE_MANUAL``) eklenir.
    """
    from api.management.commands.aggregate_readings import recompute_range
    from api.models import Reading
    from scada_io.persistence import _apply_decimals

    from .models import SimOutboxEntry
    from .services import build_sim_payloads_for_times
    from .sim_outbox import accept_window_hours

    cabinet = record.cabinet
    cols = {c["code"]: c for c in template_columns(cabinet, include_unsent=True)}
    rows = record.rows or []
    if not rows:
        raise ManualImportError("İçe aktarılacak satır yok.")

    times = [datetime.fromisoformat(r["t"]) for r in rows]
    start, end = min(times), max(times)
    sensor_ids = [c["sensor"].id for c in cols.values()]
    existing = _existing_keys(sensor_ids, start, end)
    status_cache: dict[int, object] = {}

    new_readings = []
    skipped = 0
    dropped_cols: set[str] = set()
    inserted_minutes: set[datetime] = set()
    for r, t in zip(rows, times):
        code = int(r["s"])
        if code not in status_cache:
            status_cache[code] = _status_obj(code)
        quality = status_quality(code)
        for pcode, val in (r.get("v") or {}).items():
            col = cols.get(pcode)
            if col is None:   # yükleme ile onay arasında sensör/kolon kaldırıldı
                dropped_cols.add(pcode)
                continue
            sensor = col["sensor"]
            if (sensor.id, t) in existing:
                skipped += 1
                continue
            new_readings.append(Reading(
                sensor=sensor,
                value=_apply_decimals(float(val), sensor.decimals),
                status=status_cache[code],
                quality=quality,
                origin="manual",
                time_iso=t,
            ))
            inserted_minutes.add(t)

    with transaction.atomic():
        Reading.objects.bulk_create(new_readings, batch_size=2000)
        if new_readings:
            recompute_range(
                start, end + timedelta(minutes=1),
                sensor_ids={rd.sensor_id for rd in new_readings},
            )

    # --- SİM kuyruğu (commit sonrası) ---
    cutoff = timezone.now() - timedelta(hours=accept_window_hours())
    targets = sorted(t for t in inserted_minutes if t >= cutoff)
    enqueued = 0
    by_day: dict = defaultdict(list)
    for t in targets:   # tek sorgu penceresi makul kalsın diye gün gün
        by_day[t.date()].append(t)
    for chunk in by_day.values():
        for payload in build_sim_payloads_for_times(cabinet, chunk):
            if payload.is_empty or not payload.readtime_dt:
                continue
            _, created = SimOutboxEntry.enqueue(
                cabinet,
                readtime=payload.readtime_dt,
                payload=payload.values,
                period=payload.period,
                priority=SimOutboxEntry.PRIORITY_BACKFILL,
                source=SimOutboxEntry.SOURCE_MANUAL,
            )
            enqueued += int(created)
    if enqueued:
        from .tasks import drain_sim_outbox
        transaction.on_commit(
            lambda: drain_sim_outbox.apply_async(args=[cabinet.id], expires=120)
        )

    return {
        "inserted": len(new_readings),
        "skipped": skipped,
        "enqueued": enqueued,
        "minutes": len(inserted_minutes),
        "dropped_columns": sorted(dropped_cols),
    }
