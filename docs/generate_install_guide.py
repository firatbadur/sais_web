# -*- coding: utf-8 -*-
"""Envisoft WebX kurulum kılavuzu (.docx) üretici — tek doğruluk kaynağı.

Kılavuz metni bu script'te tutulur; `docs/Envisoft-WebX-Kurulum-Kilavuzu.docx`
üretilir. PDF, docx'ten ayrı bir adımda türetilir (yerelde Word/docx2pdf,
CI'da LibreOffice `soffice --convert-to pdf`).

Kullanım:
    python docs/generate_install_guide.py [cikti.docx]

Installer / kurulum davranışı değiştiğinde bu script güncellenir; GitHub Actions
(.github/workflows/install-guide.yml) installer/** değişiminde docx+pdf'i otomatik
yeniden üretip commit'ler.
"""
import os
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BRAND = RGBColor(0x1B, 0x5E, 0x20)
ACCENT = RGBColor(0x2E, 0x7D, 0x32)
GREY = RGBColor(0x55, 0x55, 0x55)
CODE_BG = "F2F2F2"
HDR_BG = "1B5E20"
WARN_BG = "FFF4E5"
OK_BG = "E8F5E9"
ERR_BG = "FDECEA"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.15


def _shade(el, fill):
    pPr = el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def _cell_shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def h1(text):
    doc.add_page_break()
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        r.font.color.rgb = BRAND
        r.font.size = Pt(18)
    return p


def h2(text):
    p = doc.add_heading(text, level=2)
    for r in p.runs:
        r.font.color.rgb = ACCENT
        r.font.size = Pt(14)
    return p


def h3(text):
    p = doc.add_heading(text, level=3)
    for r in p.runs:
        r.font.color.rgb = ACCENT
        r.font.size = Pt(12)
    return p


def para(text="", bold=False, italic=False, color=None, size=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    return p


def bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.25 + 0.25 * level)
    for i, seg in enumerate(text.split("**")):
        p.add_run(seg).bold = (i % 2 == 1)
    return p


def numbered(text):
    p = doc.add_paragraph(style="List Number")
    for i, seg in enumerate(text.split("**")):
        p.add_run(seg).bold = (i % 2 == 1)
    return p


def code(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.1)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    _shade(p._p, CODE_BG)
    for i, ln in enumerate(text.strip("\n").split("\n")):
        if i:
            p.add_run().add_break()
        r = p.add_run(ln)
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    return p


def callout(title, text, kind="warn"):
    fill = {"warn": WARN_BG, "ok": OK_BG, "err": ERR_BG}[kind]
    icon = {"warn": "⚠  ", "ok": "✓  ", "err": "✕  "}[kind]
    tcolor = {"warn": RGBColor(0xB7, 0x4A, 0x00), "ok": BRAND, "err": RGBColor(0xC0, 0x28, 0x28)}[kind]
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    _cell_shade(cell, fill)
    p = cell.paragraphs[0]
    r = p.add_run(icon + title)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = tcolor
    cell.add_paragraph().add_run(text).font.size = Pt(10.5)
    doc.add_paragraph()
    return tbl


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        _cell_shade(hdr[i], HDR_BG)
        run = hdr[i].paragraphs[0].add_run(htext)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(10.5)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            pr = cells[i].paragraphs[0]
            for j, seg in enumerate(str(val).split("**")):
                rr = pr.add_run(seg)
                rr.bold = (j % 2 == 1)
                rr.font.size = Pt(10)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


def error_block(baslik, belirti, sebep, cozum_lines, cmd=None):
    h3(baslik)
    p = doc.add_paragraph()
    p.add_run("Belirti: ").bold = True
    p.add_run(belirti)
    p = doc.add_paragraph()
    p.add_run("Sebep: ").bold = True
    p.add_run(sebep)
    p = doc.add_paragraph()
    p.add_run("Çözüm:").bold = True
    for c in cozum_lines:
        bullet(c)
    if cmd:
        code(cmd)


# ============================ KAPAK ============================
for _ in range(3):
    doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("ENVISOFT WebX"); r.bold = True; r.font.size = Pt(40); r.font.color.rgb = BRAND
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Kurulum Kılavuzu"); r.font.size = Pt(24); r.font.color.rgb = ACCENT
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Atıksu Sürekli İzleme (SAIS) Web SCADA Uygulaması")
r.italic = True; r.font.size = Pt(13); r.font.color.rgb = GREY
for _ in range(2):
    doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Windows ve Linux Sunucu Kurulumu\n"
              "Kurulum Öncesi Kontroller · Adım Adım Kurulum · Karşılaşılabilecek Hatalar")
r.font.size = Pt(13); r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
for _ in range(6):
    doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Envisoft — Kurulum ve Devreye Alma Ekibi için")
r.font.size = Pt(11); r.font.color.rgb = GREY

# ============================ İÇİNDEKİLER ============================
h1("İçindekiler")
toc = [
    "1. Giriş",
    "2. Kurulum Öncesi Kontroller (kuruluma başlamadan ÖNCE okuyun)",
    "   2.1 Donanım ve Kaynak Gereksinimleri",
    "   2.2 Sanallaştırma Kontrolü (en kritik adım)",
    "   2.3 Windows Sürüm Kontrolü",
    "   2.4 Ağ, Port ve DNS Gereksinimleri",
    "   2.5 Kurulumda İhtiyacınız Olacak Bilgiler",
    "3. Windows Kurulumu (kurulum sihirbazı .exe)",
    "4. Linux Kurulumu (Ubuntu / Debian sunucu)",
    "5. Kurulum Sonrası: Domain + SSL (HTTPS)",
    "6. Doğrulama ve Yönetim",
    "7. Karşılaşılabilecek Hatalar ve Çözümleri",
    "   7.1 Kurulum öncesi / sırasında",
    "   7.2 İmaj indirme",
    "   7.3 Web erişimi (404 / 502 / 400)",
    "   7.4 Domain ve SSL",
    "   7.5 Hızlı başvuru tablosu",
    "8. Hızlı Kontrol Listesi",
    "Ek A. Yönetim Komutları Referansı",
]
for t in toc:
    p = doc.add_paragraph()
    run = p.add_run(t)
    if not t.startswith("   "):
        run.bold = True
    run.font.size = Pt(12)

# ============================ 1. GİRİŞ ============================
h1("1. Giriş")
para("Envisoft WebX, atıksu sürekli izleme istasyonlarından (Çevre Bakanlığı SAIS rejimi) gelen "
     "ölçüm verilerini toplayan, saklayan ve izleyen bir web tabanlı SCADA uygulamasıdır. Uygulama "
     "Docker konteyner teknolojisiyle paketlenmiştir; veritabanı, önbellek, web sunucusu ve arka plan "
     "görevleri tek bir yığın (stack) halinde çalışır.")
h2("İki kurulum yöntemi vardır")
table(
    ["Yöntem", "Ne zaman kullanılır", "Nasıl"],
    [
        ["**Windows Kurulumu**", "Sanallaştırma açık, güncel Windows (10/11 veya Server 2022+)", "Tek tıkla kurulum sihirbazı (.exe) — WSL2 + Docker'ı kendisi kurar"],
        ["**Linux Kurulumu**", "Bulut/kiralık sunucu veya WSL2 desteklemeyen Windows (ör. Server 2019)", "Tek komutluk kurulum scripti (install-linux.sh)"],
    ],
    widths=[1.6, 3.0, 2.4],
)
callout("Hangisini seçmeliyim?",
        "Elinizde sanallaştırma açık modern bir Windows varsa Windows kurulumu en kolayıdır. "
        "Kiralık/bulut sunucu (VPS) alıyorsanız veya Windows sürümünüz eskiyse (Server 2019) "
        "Linux kurulumu daha sağlıklı ve sorunsuzdur.", kind="ok")

# ============================ 2. ÖN KONTROLLER ============================
h1("2. Kurulum Öncesi Kontroller")
para("Bu bölümdeki kontroller kuruluma başlamadan ÖNCE yapılmalıdır. Özellikle 2.2 (sanallaştırma) "
     "ve 2.3 (Windows sürümü) atlanırsa kurulum yarıda başarısız olur.", bold=True)

h2("2.1 Donanım ve Kaynak Gereksinimleri")
table(
    ["Kaynak", "Minimum", "Önerilen"],
    [
        ["CPU (çekirdek)", "2", "4 veya üzeri"],
        ["RAM", "4 GB", "8 GB veya üzeri"],
        ["Disk (boş alan)", "40 GB", "80 GB veya üzeri (veri biriktikçe)"],
        ["İnternet", "Kesintisiz çıkış", "Kesintisiz çıkış (Bakanlık SIM için şart)"],
    ],
    widths=[2.2, 2.0, 3.0],
)

h2("2.2 Sanallaştırma Kontrolü (EN KRİTİK ADIM)")
para("Uygulama Docker ile çalışır. Windows'ta Docker, WSL2 üzerinden çalışır ve WSL2 donanım "
     "sanallaştırması (Intel VT-x / AMD-V) gerektirir. Bu kapalıysa kurulum ilerlemez.")
h3("A) Fiziksel bilgisayara kuruyorsanız")
numbered("Ctrl + Shift + Esc ile Görev Yöneticisi'ni açın.")
numbered("Performans sekmesi → CPU'ya tıklayın.")
numbered("Sağ altta **Sanallaştırma (Virtualization)** satırına bakın.")
table(
    ["Görünen", "Anlamı", "Yapılacak"],
    [
        ["Etkin (Enabled)", "Sanallaştırma açık", "Kuruluma devam edebilirsiniz"],
        ["Devre dışı (Disabled)", "Kapalı", "BIOS/UEFI'de VT-x / AMD-V (SVM) seçeneğini açın"],
    ],
    widths=[2.0, 2.2, 3.0],
)
h3("B) Sanal sunucuya (VM) kuruyorsanız — iç içe sanallaştırma (nested)")
para("Bir sanal makinenin içine kurulum yapıyorsanız, host (ana) hipervizörde nested virtualization "
     "açık olmalıdır.")
bullet("Görev Yöneticisi'nde \"Sanallaştırma\" yerine \"Sanal makine: Evet\" yazıyorsa, bu satır nested durumunu GÖSTERMEZ.")
bullet("**Kesin test:** yönetici PowerShell'de WSL2'yi kurmayı deneyin:")
code("wsl --install --no-distribution\n# yeniden başlatın, sonra:\nwsl --set-default-version 2\nwsl --install -d Ubuntu")
bullet("Kurulup açılıyorsa nested AÇIK. \"sanallaştırma etkin değil / 0x80370102\" hatası → nested KAPALI.")
callout("Kiralık sunucu (VPS) aldıysanız",
        "Nested virtualization ayarı sağlayıcının elindedir. Sağlayıcıya \"nested virtualization "
        "açık mı?\" diye sorun. Açamıyorlarsa (çoğu ucuz VPS'te kapalıdır) onlardan bir LINUX sunucu "
        "(Ubuntu 22.04/24.04) isteyin — Linux'ta Docker sanallaştırma gerektirmez (Bölüm 4).")

h2("2.3 Windows Sürüm Kontrolü")
para("WSL2 yalnızca güncel Windows'ta çalışır. Yönetici PowerShell'de kontrol edin:")
code("Get-ComputerInfo -Property OsName, OsVersion, OsBuildNumber")
table(
    ["Windows sürümü", "Durum", "Sonuç"],
    [
        ["Windows 10 (build 19041+)", "Destekler", "Windows kurulumu"],
        ["Windows 11", "Destekler", "Windows kurulumu"],
        ["Windows Server 2022 / 2025", "Destekler", "Windows kurulumu"],
        ["**Windows Server 2019** (build 17763)", "**Desteklemez**", "**Linux sunucu (Bölüm 4)**"],
    ],
    widths=[3.0, 1.8, 2.2],
)
callout("Windows Server 2019 uyarısı",
        "Server 2019'da WSL2 çalışmaz (build 17763 < 19041). Nested açık olsa bile Windows kurulumu "
        "başarısız olur → Ubuntu/Debian sunucuya Linux kurulumu (Bölüm 4) yapın.")

h2("2.4 Ağ, Port ve DNS Gereksinimleri")
bullet("**80 ve 443 portları** dışarıya açık olmalı (modem/firewall yönlendirmesi; bulut sunucuda sağlayıcı güvenlik duvarı).")
bullet("**Domain (alan adı):** A kaydı, sunucunun genel IP'sine yönlendirilmeli. Örn: demo.envisoft.com.tr → 62.171.190.67")
bullet("**Çıkış interneti:** imaj indirme (GHCR), SSL (Let's Encrypt), Bakanlık SIM için gerekli.")
bullet("**Cloudflare/CDN kullanacaksanız:** SSL modu \"Full (strict)\" veya kayıt \"DNS only\" olmalı (Bölüm 5.3).")

h2("2.5 Kurulumda İhtiyacınız Olacak Bilgiler")
bullet("**Domain / alan adı** (varsa)")
bullet("**Yönetici (admin) kullanıcı adı ve şifresi**")
bullet("**Yönetici e-postası** (SSL bildirimleri için)")
bullet("**Lisans anahtarı ve URL** (varsa; yoksa boş geçilebilir)")
bullet("**GHCR erişimi** — yalnız geliştirme ekibi için; müşteri paketlerinde gömülüdür, sorulmaz.")

# ============================ 3. WINDOWS ============================
h1("3. Windows Kurulumu (kurulum sihirbazı)")
para("Ön koşul (Bölüm 2): sanallaştırma açık + Windows 10 (19041+) / 11 / Server 2022+.")
numbered("**EnvisoftWebX-Setup-vX.Y.Z.exe** dosyasını sunucuya indirin.")
numbered("Sağ tık → **Yönetici olarak çalıştır**.")
numbered("Sihirbazı doldurun: **lisans modu**, **veritabanı**, **domain / SSL**, **yönetici hesabı**.")
bullet("**Lisans modu:** **Deneme** (30 gün, sonunda kilitlenir) veya **Lisanslı** (anahtar + URL girin). Lisans zorlaması güvenlik gereği yapılandırmadan kapatılamaz.")
numbered("Kurulum **WSL2 + Docker CE**'yi otomatik kurar; gerekirse **yeniden başlatır** ve açılışta kaldığı yerden devam eder.")
numbered("İmajlar indirilir, ilk veriler yüklenir, **otomatik başlatma servisi (EnvisoftWebX)** kaydedilir.")
numbered("Tarayıcıdan **http://localhost/dashboard/** ile yönetici hesabıyla giriş yapın.")
callout("Kurulum sonrası",
        "Kurulum dizini C:\\EnvisoftWebX'tir. Uygulama Windows her açıldığında otomatik başlar. "
        "Domain + SSL için Bölüm 5'e bakın.", kind="ok")

# ============================ 4. LINUX ============================
h1("4. Linux Kurulumu (Ubuntu / Debian sunucu)")
para("Bulut/kiralık sunucular ve WSL2 desteklemeyen ortamlar için önerilir. Önerilen: "
     "**Ubuntu 22.04 LTS** veya **24.04 LTS** (Debian 12 de olur).")
h2("4.1 Sunucuya bağlanın (SSH)")
code("ssh root@SUNUCU_IP\n# özel SSH portu kullanıyorsanız:\nssh -p 2222 root@SUNUCU_IP")
h2("4.2 Kurulum scriptini kopyalayın")
para("Kurulum tek dosyadır: install-linux.sh. Kendi bilgisayarınızda YENİ bir PowerShell'de:")
code("scp install-linux.sh root@SUNUCU_IP:~/\n# SSH portu farklıysa (büyük P):\nscp -P 2222 install-linux.sh root@SUNUCU_IP:~/")
bullet("Hedef olarak **~/** kullanın (/root/ yerine) — yazma izni sorunu olmaz.")
h2("4.3 Scripti çalıştırın")
code("sudo bash ~/install-linux.sh")
para("Şifre sorarsa giriş kullanıcınızın şifresini yazın (Linux'ta şifre yazarken ekranda "
     "görünmez — normaldir, yazıp Enter'a basın).")
h2("4.4 Sorulara cevap verin")
bullet("**Alan adı (domain):** SSL'i sonra ayarlayacak olsanız bile GİRİN — Django'nun tanıması için gerekli. (Boş geçerseniz sonradan elle eklemek gerekir, Bölüm 5.4.)")
bullet("**Yönetici kullanıcı adı, şifresi, e-postası**")
bullet("**Lisans modu:** [1] Lisanslı (anahtar + URL) veya [2] Deneme (30 gün, sonunda kilitlenir). Varsayılan Deneme.")
bullet("**GHCR token** — yalnız geliştirici sürümünde sorulur; müşteri paketinde gömülüdür.")
h2("4.5 Otomatik akış")
numbered("Docker CE kurulur (yoksa)")
numbered("Yapılandırma (.env) üretilir; güvenlik anahtarı + DB şifresi rastgele oluşturulur")
numbered("İmajlar indirilir (ağ koparsa otomatik tekrar dener)")
numbered("Yığın başlar, ilk veriler yüklenir, yönetici hesabı oluşturulur")
numbered("**systemd servisi** (envisoft-webx) kaydedilir → açılışta otomatik başlar")
para("Bitince ekranda dashboard adresi yazar: **http://SUNUCU_IP/dashboard/**")

# ============================ 5. DOMAIN + SSL ============================
h1("5. Kurulum Sonrası: Domain + SSL (HTTPS)")
para("Uygulama dahili Caddy sunucusuyla otomatik HTTPS (Let's Encrypt) sağlar.")
h2("5.1 DNS kaydı")
bullet("Bir **A kaydı** ekleyin: alt alan adı → sunucunun genel IP'si.")
code("nslookup demo.envisoft.com.tr\n# sunucunuzun IP'sini döndürmeli")
h2("5.2 Dashboard'dan SSL'i açın")
numbered("Dashboard → **Yönetici → Web Erişim Ayarları**")
numbered("**Etkin** aç, **Domain** yaz, **TLS = Let's Encrypt**, **e-posta** gir, kaydet.")
numbered("Caddy ilk istekte sertifikayı otomatik alır (birkaç saniye).")
h2("5.3 Cloudflare / CDN kullanıyorsanız (ÖNEMLİ)")
callout("Sonsuz yönlendirme (ERR_TOO_MANY_REDIRECTS) tuzağı",
        "Cloudflare SSL modu \"Flexible\" ise Cloudflare sunucuya HTTP ile bağlanır, sunucu HTTPS'e "
        "yönlendirir → sonsuz döngü, sayfa açılmaz.", kind="err")
bullet("**DNS only (gri bulut):** kaydın yanındaki turuncu bulutu gri yapın — Caddy kendi sertifikasını alır, doğrudan HTTPS sunar.")
bullet("**Turuncu bulut kalacaksa:** Cloudflare → SSL/TLS → mod = **\"Full (strict)\"**.")
h2("5.4 ALLOWED_HOSTS (domain'i kurulumda girmediyseniz)")
para("Dashboard'dan domain eklemek yalnız Caddy'yi yönlendirir; Django'nun kabul etmesi için "
     "yapılandırmada da tanımlı olmalı. Kurulumda domain girdiyseniz otomatik yapılmıştır. Aksi halde:")
code("cd /opt/envisoft\n"
     "sed -i '/^DJANGO_ALLOWED_HOSTS=/ s/$/,.SIZIN.DOMAIN.tr/' .env\n"
     "sed -i '/^DJANGO_CSRF_TRUSTED_ORIGINS=/ s#$#,https://*.SIZIN.DOMAIN.tr#' .env\n"
     "docker compose -f docker-compose.prod.yml --env-file .env up -d")

# ============================ 6. DOĞRULAMA ============================
h1("6. Doğrulama ve Yönetim")
h2("Container durumu (Linux)")
code("cd /opt/envisoft\ndocker compose -f docker-compose.prod.yml --env-file .env ps")
para("web, db, redis, celery_worker, celery_beat, caddy, watchtower → hepsi \"Up\" / \"healthy\" olmalı.")
h2("Loglar")
code("docker compose -f docker-compose.prod.yml --env-file .env logs -f web")
h2("Otomatik başlatma servisi")
bullet("**Linux:** systemctl status envisoft-webx")
bullet("**Windows:** Hizmetler (services.msc) → EnvisoftWebX")
callout("Yeniden başlatma testi (önerilir)",
        "Kurulumdan sonra sunucuyu bir kez yeniden başlatın (reboot) ve açılışta uygulamanın "
        "kendiliğinden geldiğini doğrulayın — gerçek elektrik kesintisi senaryosu.", kind="ok")

# ============================ 7. HATALAR ============================
h1("7. Karşılaşılabilecek Hatalar ve Çözümleri")
para("Bu bölüm sahada gerçekten karşılaşılan durumlara göre hazırlanmıştır. Her başlıkta belirti, "
     "sebep ve çözüm adımları verilmiştir.")

h2("7.1 Kurulum öncesi / sırasında")
error_block(
    "\"wsl\" komutu tanınmıyor / 0x80370102 / \"sanallaştırma etkin değil\"",
    "WSL2 kurulmuyor veya başlamıyor; kurulum bu aşamada duruyor.",
    "Donanım sanallaştırması (VT-x/AMD-V) veya sanal sunucuda nested virtualization kapalı; ya da "
    "Windows sürümü WSL2 desteklemiyor (Server 2019).",
    ["Fiziksel PC: BIOS/UEFI'de VT-x / AMD-V açın.",
     "Sanal sunucu (VM): host hipervizörde nested virtualization açtırın (kiralıksa sağlayıcıya sorun).",
     "Windows Server 2019 ise: WSL2 desteklenmez → Linux sunucuya kurulum yapın (Bölüm 4).",
     "Açamıyorsanız: bir Ubuntu 22.04/24.04 sunucu edinip Bölüm 4'ü izleyin."],
)
error_block(
    "Kurulum scripti 2. adımda sessizce duruyor (eski sürüm)",
    "Linux scripti \"Yapılandırma üretiliyor\" satırından sonra hiçbir şey demeden çıkıyor.",
    "Eski script sürümünde bir kabuk (SIGPIPE) hatası vardı.",
    ["Güncel install-linux.sh sürümünü kullanın (bu sorun giderildi).",
     "Elinizdeki dosya eskiyse güncel sürümü tekrar kopyalayıp çalıştırın."],
)

h2("7.2 İmaj indirme")
error_block(
    "\"connection reset by peer\" / imaj indirme yarıda kopuyor",
    "docker pull sırasında bağlantı kopuyor (özellikle IPv6 üzerinde).",
    "Geçici ağ sorunu. İndirilen katmanlar önbelleğe alındığından tekrar denemek kaldığı yerden sürer.",
    ["Güncel script bunu otomatik 5 kez dener.",
     "Elle tekrar denemek isterseniz aşağıdaki komutu birkaç kez çalıştırın:"],
    "cd /opt/envisoft\nfor i in 1 2 3 4 5; do docker compose -f docker-compose.prod.yml --env-file .env pull && break; sleep 4; done",
)

h2("7.3 Web erişimi (404 / 502 / 400)")
error_block(
    "Dashboard'da 404 (sayfa bulunamadı)",
    "http://SUNUCU_IP/dashboard/ açılınca 404 dönüyor.",
    "Caddy, web'in ürettiği yapılandırmayı yüklemeden başlamış; stok yapılandırmayı sunuyor.",
    ["Caddy'yi bir kez yeniden başlatın (güncel script bunu kurulum sonunda otomatik yapar):"],
    "cd /opt/envisoft\ndocker compose -f docker-compose.prod.yml --env-file .env restart caddy",
)
error_block(
    "502 Bad Gateway",
    "Sayfa 502 hatası veriyor.",
    "Web container henüz başlıyor (migration/collectstatic) veya çökmüş.",
    ["Durumu ve logları kontrol edin; web yeni başladıysa 30-60 sn bekleyin:"],
    "cd /opt/envisoft\ndocker compose -f docker-compose.prod.yml --env-file .env ps\ndocker compose -f docker-compose.prod.yml --env-file .env logs --tail=60 web",
)
error_block(
    "400 Bad Request (DisallowedHost)",
    "Domain ile açınca 400 hatası; IP ile sorun yok.",
    "Domain, Django ALLOWED_HOSTS listesinde yok (dashboard'dan domain eklemek bunu güncellemez).",
    ["Bölüm 5.4'teki .env düzenlemesini yapıp web'i yeniden başlatın."],
)

h2("7.4 Domain ve SSL")
error_block(
    "Sonsuz yönlendirme (ERR_TOO_MANY_REDIRECTS)",
    "Domain açılınca tarayıcı \"çok fazla yönlendirme\" diyor.",
    "Cloudflare SSL modu \"Flexible\": Cloudflare origin'e HTTP ile gelir, sunucu HTTPS'e yönlendirir → döngü.",
    ["Cloudflare kaydını \"DNS only\" (gri bulut) yapın, VEYA",
     "Cloudflare → SSL/TLS → modu \"Full (strict)\" yapın (Bölüm 5.3)."],
)
error_block(
    "SSL sertifikası alınamıyor",
    "Caddy loglarında sertifika hatası; https açılmıyor.",
    "DNS henüz doğru IP'ye yayılmamış, 80/443 dışarı kapalı, veya Let's Encrypt hız limiti.",
    ["nslookup ile domain'in sunucu IP'sini döndürdüğünü doğrulayın.",
     "80 ve 443'ün dışarıdan erişilebilir olduğunu kontrol edin.",
     "Caddy loglarına bakın:"],
    "cd /opt/envisoft\ndocker compose -f docker-compose.prod.yml --env-file .env logs -f caddy",
)

h2("7.5 Hızlı başvuru tablosu")
table(
    ["Belirti", "Olası sebep", "Çözüm"],
    [
        ["wsl tanınmıyor / 0x80370102", "Sanallaştırma / nested kapalı", "VT-x/AMD-V veya nested aç; olmuyorsa Linux sunucu"],
        ["404 dashboard", "Caddy config yüklenmemiş", "restart caddy"],
        ["502 Bad Gateway", "Web başlıyor / çökmüş", "ps + logs web; 30-60 sn bekle"],
        ["400 DisallowedHost", "Domain ALLOWED_HOSTS'ta yok", ".env'e ekle + up -d (Bölüm 5.4)"],
        ["Sonsuz yönlendirme", "Cloudflare Flexible", "DNS only veya Full (strict)"],
        ["İmaj indirme koptu", "Geçici ağ / IPv6", "pull'u tekrarla (önbellekten sürer)"],
        ["Sertifika alınamadı", "DNS/port/rate-limit", "DNS + 80/443 + caddy log"],
    ],
    widths=[2.3, 2.0, 3.0],
)

# ============================ 8. KONTROL LİSTESİ ============================
h1("8. Hızlı Kontrol Listesi")
h2("Kurulum öncesi")
for t in [
    "Sanallaştırma açık mı? (Windows: Görev Yöneticisi → CPU → Sanallaştırma: Etkin)",
    "VM ise nested virtualization açık mı? (WSL2 ile test edildi)",
    "Windows sürümü uygun mu? (10/11 veya Server 2022+; 2019 DEĞİL)",
    "En az 4 GB RAM, 2 çekirdek, 40 GB boş disk var mı?",
    "80 ve 443 portları dışarıya açık mı?",
    "Domain A kaydı sunucu IP'sine yönlendirildi mi?",
    "Yönetici bilgileri ve (varsa) domain/lisans hazır mı?",
]:
    p = doc.add_paragraph(); p.add_run("☐  ").bold = True; p.add_run(t)
h2("Kurulum sonrası")
for t in [
    "Tüm container'lar Up / healthy mi?",
    "http://SUNUCU_IP/dashboard/ ile giriş yapılıyor mu?",
    "Domain + SSL açıldı, https:// ile açılıyor mu?",
    "Cloudflare varsa mod \"Full (strict)\" / \"DNS only\" mu?",
    "Yeniden başlatma sonrası uygulama otomatik geldi mi?",
    "Station / Connection / Sensor / SaisCabinet kayıtları girildi mi?",
]:
    p = doc.add_paragraph(); p.add_run("☐  ").bold = True; p.add_run(t)

# ============================ EK A ============================
h1("Ek A. Yönetim Komutları Referansı (Linux)")
para("Komutlar sunucuda /opt/envisoft dizininde çalıştırılır.")
code("cd /opt/envisoft\nC='docker compose -f docker-compose.prod.yml --env-file .env'\n\n"
     "$C ps                 # servis durumu\n$C logs -f web        # web logları\n"
     "$C restart caddy      # Caddy yeniden başlat (404 çözümü)\n"
     "$C up -d              # yığını başlat / güncel .env uygula\n$C down               # yığını durdur")
code("systemctl status envisoft-webx     # otomatik başlatma servisi\n"
     "systemctl restart envisoft-webx")
para("Sürüm güncelleme dashboard'daki \"Şimdi Güncelle\" butonuyla yapılır; rollback için .env "
     "içinde IMAGE_TAG değiştirilip \"up -d\" çalıştırılır.")

doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("— Envisoft WebX Kurulum Kılavuzu —"); r.italic = True; r.font.color.rgb = GREY

# --- Kaydet ---
default_out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                           "Envisoft-WebX-Kurulum-Kilavuzu.docx")
out = sys.argv[1] if len(sys.argv) > 1 else default_out
doc.save(out)
print("KAYDEDILDI:", out)
